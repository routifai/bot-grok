"""Live sheets: a CSV / XLSX kept up to date cell by cell, by the same rules as a live page."""

from __future__ import annotations

import io
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import pytest

from omnigent.errors import OmnigentError
from omnigent.stores.artifact_store.local import LocalArtifactStore
from omnigent.stores.scheduled_task_store.sqlalchemy_store import SqlAlchemyScheduledTaskStore
from omnigent.superchat.artifact_kinds import KIND_MIME
from omnigent.superchat.artifacts.parts import PART_EDITORS, part_editor
from omnigent.superchat.artifacts.store import SqlAlchemyArtifactStore
from omnigent.superchat.live.service import LiveService, task_prompt
from omnigent.superchat.live.store import SqlAlchemyLiveStore
from omnigent.superchat.sheets import FEATURE  # noqa: F401  (registers the cell editor)
from tests.superchat.test_live_artifacts import _SESSION, _Convs, _Scheduler

CSV = b"tea,price\nSencha,12.50\nMatcha,$30\n"


def _xlsx() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Prices"
    for row in (["tea", "price"], ["Sencha", 12.5], ["Matcha", 30], ["Total", "=SUM(B2:B3)"]):
        ws.append(row)
    wb.create_sheet("Notes").append(["hello"])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


@pytest.fixture()
def env(db_uri: str, tmp_path: Path):
    artifacts = SqlAlchemyArtifactStore(db_uri, lambda: LocalArtifactStore(str(tmp_path / "b")))
    service = LiveService(
        artifacts, SqlAlchemyLiveStore(db_uri), SqlAlchemyScheduledTaskStore(db_uri), _Convs(),
        lambda: _Scheduler(),
    )  # fmt: skip
    service.loop = SimpleNamespace(call_soon_threadsafe=lambda call, *args: call(*args))

    def make(name: str, data: bytes) -> Any:
        kind = name.rsplit(".", 1)[1]
        item = artifacts.create(
            user_id=None, parent_session_id=_SESSION, name=name, title="Teas", kind=kind,
            mime=KIND_MIME[kind], data=data,
        )  # fmt: skip
        service.upsert(item, rrule="FREQ=DAILY;BYHOUR=8;BYMINUTE=0", sources=["teas.json"])
        return item

    return SimpleNamespace(artifacts=artifacts, service=service, make=make)


def _save(env: Any, item: Any, changes: list[dict[str, Any]]) -> dict[str, Any]:
    head = env.artifacts.newest(item)
    return env.service.save(
        head, base_version=head.version, changes=changes, facts={}, sources=None
    )


def _bytes(env: Any, item: Any) -> bytes:
    return env.artifacts.read(env.artifacts.newest(item))


def test_a_csv_is_updated_cell_by_cell_and_a_new_row_is_added(env) -> None:
    item = env.make("teas.csv", CSV)
    ctx = env.service.context(item)
    assert ctx["source"].splitlines()[1] == "A2: Sencha | B2: 12.50" and "B3" in ctx["ids"]
    out = _save(
        env,
        item,
        [
            {"id": "B2", "after": "13.00"},
            {"id": "B3", "after": "$30"},  # same as the sheet shows: nothing to do
            {"id": "A4", "after": "Oolong"},
            {"id": "Sheet1!B4", "after": "18.25"},
        ],
    )
    assert [(c["id"], c["before"], c["after"]) for c in out["applied"]] == [
        ("B2", "12.50", "13.00"),
        ("A4", "", "Oolong"),
        ("Sheet1!B4", "", "18.25"),
    ]
    assert _bytes(env, item) == b"tea,price\nSencha,13.00\nMatcha,$30\nOolong,18.25\n"
    head = env.artifacts.newest(item)
    assert (head.version, head.origin, head.edit_summary) == (2, "update", "3 changes")
    # 13 and 13.00 are the same number: a run that reads it differently changes nothing.
    assert _save(env, item, [{"id": "B2", "after": "13"}])["version"] is None


def test_a_cell_the_person_changed_becomes_a_suggestion_and_undo_restores(env) -> None:
    item = env.make("teas.csv", CSV)
    _save(env, item, [{"id": "B2", "after": "13.00"}])
    head = env.artifacts.newest(item)
    env.artifacts.create(
        user_id=None, parent_session_id=_SESSION, name=head.name, title=head.title, kind="csv",
        mime=head.mime, data=_bytes(env, item).replace(b"13.00", b"11"), origin="manual",
        parent_version_id=head.id,
    )  # fmt: skip
    out = _save(env, item, [{"id": "B2", "after": "14.00"}, {"id": "A2", "after": "Sencha A"}])
    assert out["suggestions"] == [
        {"id": "B2", "yours": "11", "source_value": "14.00", "source": None}
    ]
    assert b"Sencha A,11\n" in _bytes(env, item)
    env.service.undo(env.artifacts.newest(item))
    assert b"Sencha,11\n" in _bytes(env, item)


def test_an_xlsx_keeps_numbers_as_numbers_and_never_overwrites_a_formula(env) -> None:
    from openpyxl import load_workbook

    item = env.make("teas.xlsx", _xlsx())
    source = env.service.context(item)["source"]
    assert "Prices!B2: 12.5" in source and "Prices!B4: =SUM(B2:B3)" in source
    assert "Notes!A1: hello" in source
    out = _save(
        env,
        item,
        [
            {"id": "Prices!B2", "after": "13.50"},
            {"id": "Prices!B4", "after": "99"},  # a formula follows its inputs
            {"id": "Nope!A1", "after": "x"},
            {"id": "B2:B3", "after": "x"},
        ],
    )
    assert [c["id"] for c in out["applied"]] == ["Prices!B2"]
    assert out["unplaced"] == ["Prices!B4", "Nope!A1", "B2:B3"]
    ws = load_workbook(io.BytesIO(_bytes(env, item)))["Prices"]
    assert ws["B2"].value == 13.5 and ws["B4"].value == "=SUM(B2:B3)"
    assert _save(env, item, [{"id": "Prices!B2", "after": "13.5"}])["version"] is None


def test_only_pages_and_sheets_can_be_live_and_the_run_is_told_which(env) -> None:
    assert set(PART_EDITORS) >= {"html", "csv", "xlsx"}
    assert part_editor("html", "talk.deck.html") is None and part_editor("pdf", "a.pdf") is None
    assert 'the live sheet "teas.csv"' in task_prompt("teas.csv", "csv")
    assert 'the live page "p.html"' in task_prompt("p.html")
    doc = env.artifacts.create(
        user_id=None, parent_session_id=_SESSION, name="n.md", title="N", kind="md",
        mime=KIND_MIME["md"], data=b"# n",
    )  # fmt: skip
    env.service.upsert(doc, rrule="FREQ=DAILY;BYHOUR=8;BYMINUTE=0")  # the route refuses this
    with pytest.raises(OmnigentError):
        env.service.context(doc)
