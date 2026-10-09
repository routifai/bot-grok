"""A spreadsheet or CSV file as a schema, never as rows of text.

Numbers in a table are answered by code (pandas, duckdb) in the Computer, not by reading rows as
text, so the index holds no row of a CSV or workbook. What it keeps instead:

* a **profile** per sheet: the columns with an inferred type, the row count and the first few
  rows. It is the manifest a model reads (``files_get``, an attachment's context), written as
  Markdown with one page per sheet;
* a **schema card**: the file name, its sheet names and column names, as the one passage the
  search can find ("which file has a ``revenue`` column?").

Pure apart from reading the bytes; nothing here calls a model.
"""

from __future__ import annotations

import csv
import io
import re
from collections.abc import Iterable, Iterator
from dataclasses import dataclass, field
from datetime import date, datetime

from omnigent.runner.knowledge.chunking import MAX_CHARS, Block, Chunk, ParsedPage
from omnigent.runner.knowledge.markdown import to_markdown

#: File kinds handled as tables.
TABULAR_KINDS = frozenset({"csv", "xlsx"})
#: The chunk kind the schema card is stored with.
SCHEMA_CHUNK_KIND = "schema"
PREVIEW_ROWS = 5
#: Rows read to infer a column's type.
TYPE_SAMPLE_ROWS = 1000
MAX_SHEETS = 10
MAX_COLUMNS = 60
#: Columns shown in the preview table, and the longest cell shown.
PREVIEW_COLUMNS = 20
CELL_CHARS = 30

_INTEGER = re.compile(r"^[+-]?\d+$")
_NUMBER = re.compile(r"^[+-]?(\d+\.?\d*|\.\d+)([eE][+-]?\d+)?$")
_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}([ T]\d{2}:\d{2}(:\d{2})?)?$")
_BOOLEAN = frozenset({"true", "false"})


@dataclass(frozen=True)
class Column:
    """One column: its header text and the type its values share."""

    name: str
    type: str


@dataclass(frozen=True)
class SheetProfile:
    """One sheet (a CSV file is one unnamed sheet)."""

    name: str
    columns: list[Column]
    rows: int
    preview: list[list[str]] = field(default_factory=list)


@dataclass(frozen=True)
class TableProfile:
    """What a table file holds, without its rows."""

    sheets: list[SheetProfile]


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    return str(value).strip()


def _shown(value: object) -> str:
    text = _text(value).replace("\n", " ")
    return text if len(text) <= CELL_CHARS else text[: CELL_CHARS - 1] + "…"


def _type_of(value: object) -> str | None:
    """The type of one cell value (``None`` for an empty cell)."""
    if value is None:
        return None
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, int):
        return "integer"
    if isinstance(value, float):
        return "integer" if value.is_integer() else "number"
    if isinstance(value, datetime | date):
        return "date"
    text = str(value).strip()
    if not text:
        return None
    if _INTEGER.match(text):
        return "integer"
    if _NUMBER.match(text):
        return "number"
    if _DATE.match(text):
        return "date"
    if text.lower() in _BOOLEAN:
        return "boolean"
    return "text"


def _merge(seen: set[str]) -> str:
    """One column type from the types of its sampled values."""
    if not seen:
        return "empty"
    if len(seen) == 1:
        return next(iter(seen))
    if seen == {"integer", "number"}:
        return "number"
    return "text"


def _profile_rows(name: str, rows: Iterable[list[object]]) -> SheetProfile:
    """A sheet from its non-empty rows, header first."""
    iterator: Iterator[list[object]] = iter(rows)
    header = next(iterator, None)
    if header is None:
        return SheetProfile(name, [], 0)
    names = [_text(h) or f"column {i + 1}" for i, h in enumerate(header[:MAX_COLUMNS])]
    seen: list[set[str]] = [set() for _ in names]
    preview: list[list[str]] = []
    count = 0
    for row in iterator:
        count += 1
        if count <= TYPE_SAMPLE_ROWS:
            for index, kind_set in enumerate(seen):
                kind = _type_of(row[index]) if index < len(row) else None
                if kind:
                    kind_set.add(kind)
        if len(preview) < PREVIEW_ROWS:
            preview.append([_shown(row[i]) if i < len(row) else "" for i in range(len(names))])
    columns = [Column(n, _merge(s)) for n, s in zip(names, seen, strict=True)]
    return SheetProfile(name, columns, count, preview)


def _csv_rows(text: str) -> Iterator[list[object]]:
    sample = text[:4096]
    try:
        dialect: type[csv.Dialect] = csv.Sniffer().sniff(sample, delimiters=",;\t|")  # type: ignore[assignment]
    except csv.Error:
        dialect = csv.excel
    for row in csv.reader(io.StringIO(text), dialect):
        if any(c.strip() for c in row):
            yield list(row)


def _xlsx_profile(data: bytes) -> TableProfile:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl raises many types for a bad file
        raise ValueError(f"not a readable workbook ({type(exc).__name__})") from None
    try:
        sheets = [
            _profile_rows(
                sheet.title,
                (list(r) for r in sheet.iter_rows(values_only=True) if any(_text(c) for c in r)),
            )
            for sheet in workbook.worksheets[:MAX_SHEETS]
        ]
    finally:
        workbook.close()
    return TableProfile(sheets or [SheetProfile("Sheet1", [], 0)])


def profile_table(kind: str, data: bytes, *, name: str = "") -> TableProfile:
    """The profile of a CSV or XLSX file.

    :raises ValueError: the bytes cannot be read as that kind.
    """
    if kind == "xlsx":
        return _xlsx_profile(data)
    if kind == "csv":
        from omnigent.runner.knowledge.parse import decode_text

        return TableProfile([_profile_rows(name or "data", _csv_rows(decode_text(data)))])
    raise ValueError(f"{kind} is not a table")


def manifest_pages(profile: TableProfile, file_name: str) -> list[ParsedPage]:
    """The manifest as pages (one per sheet) so it is written with the usual page markers."""
    pages: list[ParsedPage] = []
    for number, sheet in enumerate(profile.sheets, start=1):
        blocks = [Block(sheet.name, heading_level=1)]
        if number == 1:
            blocks.append(
                Block(
                    f"Table file {file_name}: {len(profile.sheets)} sheet(s). The rows are not "
                    "indexed or shown as text; compute on the file with code (pandas, duckdb) "
                    "in the Computer."
                )
            )
        if not sheet.columns:
            blocks.append(Block("This sheet is empty."))
        else:
            cols = ", ".join(f"{c.name} ({c.type})" for c in sheet.columns)
            blocks.append(Block(f"Rows: {sheet.rows}. Columns: {cols}"))
            shown = sheet.columns[:PREVIEW_COLUMNS]
            header = " | ".join(c.name for c in shown)
            lines = [header, *(" | ".join(row[: len(shown)]) for row in sheet.preview)]
            blocks.append(Block("\n".join(lines), atomic=True))
        pages.append(ParsedPage(number, blocks))
    return pages


def manifest_markdown(profile: TableProfile, file_name: str) -> str:
    """The manifest a model reads in place of the file's rows."""
    return to_markdown(manifest_pages(profile, file_name), kind="sheet")


def schema_card(profile: TableProfile, file_name: str) -> Chunk:
    """The one passage the index keeps: file name, sheet names, column names."""
    lines = [f"Table file: {file_name}", "Sheets: " + ", ".join(s.name for s in profile.sheets)]
    lines += [
        f"Sheet {s.name} columns: " + ", ".join(c.name for c in s.columns) for s in profile.sheets
    ]
    return Chunk(1, 1, SCHEMA_CHUNK_KIND, "", "\n".join(lines)[:MAX_CHARS])
