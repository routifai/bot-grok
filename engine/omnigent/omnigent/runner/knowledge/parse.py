"""Turn a file's bytes into pages of blocks, by kind.

A PDF keeps its real pages. A text, Markdown or CSV file gets virtual pages of about
:data:`VIRTUAL_PAGE_CHARS` so a citation can still say "page 3"; an XLSX file has one page per
sheet. Nothing here calls a model.
"""

from __future__ import annotations

import csv
import io
import re
from collections.abc import Iterable, Iterator

from omnigent.runner.knowledge.chunking import TARGET_CHARS, Block, ParsedPage
from omnigent.runner.knowledge.pdf import extract_pages

#: Artifact kinds this capability indexes.
INDEXABLE_KINDS = frozenset({"pdf", "txt", "md", "csv", "xlsx"})
#: A virtual page of a text file ends at the first block boundary past this many characters.
VIRTUAL_PAGE_CHARS = 3500

_FENCE = re.compile(r"^\s*(```|~~~)")
_ATX_HEADING = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
_BLANK_RUN = re.compile(r"\n\s*\n")


def decode_text(data: bytes) -> str:
    """UTF-8 (with or without BOM); Windows-1252 for a file that is not UTF-8."""
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def paginate(blocks: Iterable[Block], limit: int = VIRTUAL_PAGE_CHARS) -> list[ParsedPage]:
    """Group blocks into virtual pages, cutting at the first block boundary past *limit*."""
    pages: list[ParsedPage] = []
    current: list[Block] = []
    size = 0
    for block in blocks:
        current.append(block)
        size += len(block.text)
        if size >= limit:
            pages.append(ParsedPage(len(pages) + 1, current))
            current, size = [], 0
    if current or not pages:
        pages.append(ParsedPage(len(pages) + 1, current))
    return pages


def _text_blocks(text: str) -> list[Block]:
    return [Block(p.strip()) for p in _BLANK_RUN.split(text.replace("\r\n", "\n")) if p.strip()]


def _markdown_blocks(text: str) -> list[Block]:
    """Headings (``#``) become heading blocks; the rest is split on blank lines, fences intact."""
    blocks: list[Block] = []
    paragraph: list[str] = []
    fenced = False

    def end_paragraph() -> None:
        body = "\n".join(paragraph).strip()
        if body:
            blocks.append(Block(body))
        paragraph.clear()

    for line in text.replace("\r\n", "\n").split("\n"):
        if _FENCE.match(line):
            fenced = not fenced
            paragraph.append(line)
            continue
        heading = None if fenced else _ATX_HEADING.match(line)
        if heading and heading.group(2):
            end_paragraph()
            blocks.append(Block(heading.group(2), heading_level=len(heading.group(1))))
        elif not fenced and not line.strip():
            end_paragraph()
        else:
            paragraph.append(line)
    end_paragraph()
    return blocks


def _row_text(cells: Iterable[object]) -> str:
    return " | ".join(_cell(c) for c in cells)


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _table_blocks(header: str, rows: Iterable[str]) -> Iterator[Block]:
    """Row slices of about :data:`TARGET_CHARS`, each starting with the header row."""
    batch: list[str] = []
    size = len(header)
    for row in rows:
        if batch and size + len(row) + 1 > TARGET_CHARS:
            yield Block("\n".join([header, *batch]), atomic=True)
            batch, size = [], len(header)
        batch.append(row)
        size += len(row) + 1
    if batch:
        yield Block("\n".join([header, *batch]), atomic=True)


def _csv_blocks(text: str) -> list[Block]:
    sample = text[:4096]
    try:
        dialect: type[csv.Dialect] = csv.Sniffer().sniff(sample, delimiters=",;\t|")  # type: ignore[assignment]
    except csv.Error:
        dialect = csv.excel
    rows = [
        _row_text(r) for r in csv.reader(io.StringIO(text), dialect) if any(c.strip() for c in r)
    ]
    if not rows:
        return []
    return list(_table_blocks(rows[0], rows[1:]))


def _xlsx_pages(data: bytes) -> list[ParsedPage]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl raises many types for a bad file
        raise ValueError(f"not a readable workbook ({type(exc).__name__})") from None
    pages: list[ParsedPage] = []
    try:
        for sheet in workbook.worksheets:
            rows = (
                _row_text(r) for r in sheet.iter_rows(values_only=True) if any(_cell(c) for c in r)
            )
            first = next(rows, None)
            blocks = [Block(sheet.title, heading_level=1)]
            if first is not None:
                blocks.extend(_table_blocks(first, rows))
            pages.append(ParsedPage(len(pages) + 1, blocks))
    finally:
        workbook.close()
    return pages or [ParsedPage(1, [])]


def parse_file(kind: str, data: bytes) -> tuple[list[ParsedPage], str]:
    """Pages of blocks for a file, and the chunk kind to store (``text``, ``csv``, ``sheet``).

    :raises ValueError: the bytes cannot be read as that kind.
    """
    if kind == "pdf":
        return extract_pages(data), "text"
    if kind == "xlsx":
        return _xlsx_pages(data), "sheet"
    text = decode_text(data)
    if kind == "csv":
        return paginate(_csv_blocks(text)), "csv"
    if kind == "md":
        return paginate(_markdown_blocks(text)), "text"
    if kind == "txt":
        return paginate(_text_blocks(text)), "text"
    raise ValueError(f"cannot index a {kind} file")
